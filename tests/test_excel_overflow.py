"""Floating totals row: two blank rows below the items, aux sections follow.

The bundled style template carries a styled data region of rows 13..60, a
totals row at 61, and the auxiliary sections (``DETAILS AS PER CHALLANS`` /
``DETAILS AS PER TALLY`` at row 71, the C&F block, and ``CLEARANCE AND
FORWARDING INVOICE`` at row 94). The generator floats the totals row so it
always sits two blank rows below the last line item: unused blank rows are
removed for short documents (no bank of empty rows) and rows are inserted for
long ones, with the aux sections moving with the totals either way.

These tests verify:

- the item rows remain dense from row 13 (Sr. no. 1..N) with no gaps;
- the totals row sits two blank rows below the last item and carries the
  invoice-amount total (e.g. 18 items in rows 13..30 put the totals at 33);
- every auxiliary label sits strictly below the totals row, so no label cell is
  overwritten by an item row and no item row is overwritten by a label;
- the auxiliary section merges follow the labels to their shifted positions.
"""

from __future__ import annotations

from boe_converter.excel_writer import (
    AUX_LABELS,
    AUX_TOTAL_USD_CELL,
    COL_AMOUNT,
    COL_SR_NO,
    ITEM_TABLE_FIRST_DATA_ROW,
    TOTALS_ROW,
    ExcelGenerator,
    _shift_coord,
    _totals_row_for,
)
from boe_converter.models import (
    ComputedDocument,
    ComputedLine,
    HeaderBlock,
    LineItem,
    RawValue,
    ReviewFlagSet,
    Totals,
)
from openpyxl.utils.cell import coordinate_to_tuple


def _raw(value) -> RawValue:
    return RawValue(raw_text=str(value), parsed=value)


def _header() -> HeaderBlock:
    return HeaderBlock(
        company_name="Gemini Unicom LLP",
        party_name=_raw("ACME"),
        usd_rate=90.0,
        details=_raw("CO-32"),
        invoice_no=_raw("INV-1"),
        invoice_date=_raw("01/01/2026"),
        be_no=_raw("BE-1"),
        be_date=_raw("01/01/2026"),
        bl_no=_raw("BL-1"),
        bl_date=_raw("01/01/2026"),
        invoice_amount=_raw(100.0),
        invoice_currency=_raw("USD"),
        package_count=_raw(1357),
        container_details=_raw("CONT-1"),
    )


def _line(serial: int) -> ComputedLine:
    item = LineItem(
        item_serial=serial,
        cth_hsn=_raw("1234"),
        description=_raw(f"ITEM-{serial}"),
        unit_price_usd=_raw(2.0),
        quantity=_raw(3.0),
        unit=_raw("PCS"),
        assessable_value=_raw(10.0),
        bcd_rate=_raw(0.1),
        bcd_amount=_raw(1.0),
        igst_rate=_raw(0.18),
        total_duty=_raw(2.0),
    )
    return ComputedLine(source=item)


def _build(n_items: int):
    lines = [_line(s) for s in range(1, n_items + 1)]
    calc_lines = lines  # ComputedLine values are not needed for positioning here
    doc = ComputedDocument(
        header=_header(),
        lines=calc_lines,
        totals=Totals(total_amount_usd=float(n_items)),
    )
    wb = ExcelGenerator(use_formulas=False).build_workbook(doc, ReviewFlagSet())
    return wb["Sheet1"], doc


def test_overflow_shifts_totals_and_aux_below_items():
    """80 items: totals + aux move below the item block, no overlap."""
    n = 80
    ws, _ = _build(n)
    first = ITEM_TABLE_FIRST_DATA_ROW

    # Dense Sr. no. 1..N from row 13, with two blank rows after the block.
    sr_values = [ws.cell(row=first + i, column=COL_SR_NO).value for i in range(n)]
    assert sr_values == list(range(1, n + 1))
    last_item_row = first + n - 1  # 92
    assert ws.cell(row=last_item_row + 1, column=COL_SR_NO).value is None
    assert ws.cell(row=last_item_row + 2, column=COL_SR_NO).value is None

    # Totals row floats to two blanks below the last item.
    shift = _totals_row_for(n) - TOTALS_ROW
    expected_totals_row = TOTALS_ROW + shift
    assert expected_totals_row == last_item_row + 3
    amount_total = ws.cell(row=expected_totals_row, column=COL_AMOUNT).value
    assert isinstance(amount_total, (int, float)) and amount_total > 0

    # Every auxiliary label sits strictly below the totals row (no collision
    # with any item row) and matches the verbatim template label at its
    # shifted coordinate.
    for coord, label in AUX_LABELS.items():
        shifted = _shift_coord(coord, shift)
        row, _col = coordinate_to_tuple(shifted)
        assert row > expected_totals_row, f"{coord}->{shifted} overlaps items/totals"
        assert ws[shifted].value == label

    # The 'total usd' value cell tracks the shifted totals amount.
    assert ws[_shift_coord(AUX_TOTAL_USD_CELL, shift)].value == amount_total

    # No item row (13..last_item_row) accidentally carries an aux label.
    aux_labels = set(AUX_LABELS.values())
    for r in range(first, last_item_row + 1):
        for c in range(1, 26):
            assert ws.cell(row=r, column=c).value not in aux_labels

    # The shifted aux section merges exist at their new positions.
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert _shift_coord("G71", shift) + ":" + _shift_coord("O71", shift) in merged
    assert _shift_coord("Q71", shift) + ":" + _shift_coord("X71", shift) in merged
    assert _shift_coord("G94", shift) + ":" + _shift_coord("O94", shift) in merged


def test_short_document_floats_totals_two_blanks_below_items():
    """18 items: totals at row 33 (not 61), aux shifted up with it."""
    n = 18
    ws, _ = _build(n)
    first = ITEM_TABLE_FIRST_DATA_ROW
    last_item_row = first + n - 1  # 30

    expected_totals_row = _totals_row_for(n)
    assert expected_totals_row == 33 == last_item_row + 3
    # The two rows between the items and the totals stay blank.
    assert ws.cell(row=last_item_row + 1, column=COL_SR_NO).value is None
    assert ws.cell(row=last_item_row + 2, column=COL_SR_NO).value is None
    amount_total = ws.cell(row=expected_totals_row, column=COL_AMOUNT).value
    assert isinstance(amount_total, (int, float)) and amount_total > 0

    # Aux labels sit below the floating totals at their shifted coordinates.
    shift = expected_totals_row - TOTALS_ROW
    for coord, label in AUX_LABELS.items():
        shifted = _shift_coord(coord, shift)
        row, _col = coordinate_to_tuple(shifted)
        assert row > expected_totals_row, f"{coord}->{shifted} overlaps items/totals"
        assert ws[shifted].value == label
    assert ws[_shift_coord(AUX_TOTAL_USD_CELL, shift)].value == amount_total

    # The shifted aux section merges exist at their new positions.
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert _shift_coord("G71", shift) + ":" + _shift_coord("O71", shift) in merged
    assert _shift_coord("Q71", shift) + ":" + _shift_coord("X71", shift) in merged
    assert _shift_coord("G94", shift) + ":" + _shift_coord("O94", shift) in merged


def test_no_bank_of_empty_rows_between_items_and_totals():
    """10 items: totals at row 25; no unused blank data rows remain."""
    ws, _ = _build(10)
    expected_totals_row = _totals_row_for(10)
    assert expected_totals_row == 25

    amount_total = ws.cell(row=expected_totals_row, column=COL_AMOUNT).value
    assert isinstance(amount_total, (int, float)) and amount_total > 0

    # Aux labels follow the totals up (shift is negative here).
    shift = expected_totals_row - TOTALS_ROW
    for coord, label in AUX_LABELS.items():
        assert ws[_shift_coord(coord, shift)].value == label

    # Original aux merges moved with the labels.
    merged = {str(rng) for rng in ws.merged_cells.ranges}
    assert _shift_coord("G71", shift) + ":" + _shift_coord("O71", shift) in merged
    assert _shift_coord("Q71", shift) + ":" + _shift_coord("X71", shift) in merged
    assert _shift_coord("G94", shift) + ":" + _shift_coord("O94", shift) in merged


def test_boundary_forty_six_items_keeps_sample_positions():
    """46 items fill rows 13..58; totals land exactly on row 61 (shift 0)."""
    n = 46
    ws, _ = _build(n)
    first = ITEM_TABLE_FIRST_DATA_ROW
    assert ws.cell(row=first + n - 1, column=COL_SR_NO).value == n  # row 58
    assert _totals_row_for(n) == TOTALS_ROW == 61
    assert ws.cell(row=TOTALS_ROW, column=COL_SR_NO).value is None  # blanks then totals
    assert ws.cell(row=TOTALS_ROW, column=COL_AMOUNT).value > 0


def _style_sig(ws, row: int, col: int):
    """A comparable per-cell style signature (bold, fill colour, borders)."""
    cell = ws.cell(row=row, column=col)
    fill = cell.fill.fgColor.rgb if (cell.fill and cell.fill.patternType) else "none"
    border = "".join(
        (getattr(cell.border, side).style or "-")[0]
        for side in ("left", "right", "top", "bottom")
    )
    return (bool(cell.font.bold), fill, border, cell.number_format)


def test_data_rows_share_uniform_styling():
    """Every item row (incl. the template's hand-edited rows 58-60 and the
    overflow continuation rows) matches the canonical data-row styling.

    The bundled template carried stray bold / missing-highlight quirks on rows
    58-60; the generator normalizes the styled data region so no item row looks
    different from the others.
    """
    n = 80  # fills rows 13..92, exercising rows 58-60 and the overflow rows
    ws, _ = _build(n)
    first = ITEM_TABLE_FIRST_DATA_ROW
    # Canonical reference: an ordinary interior data row.
    reference = {col: _style_sig(ws, first + 1, col) for col in range(1, 26)}

    for offset in range(n):
        row = first + offset
        for col in range(1, 26):
            assert _style_sig(ws, row, col) == reference[col], (
                f"row {row} col {col} style differs from the canonical data row"
            )
