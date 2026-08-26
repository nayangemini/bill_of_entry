"""BUG-002/003/004 (reopened): identity must survive the Excel round-trip.

The reported repro is *upload an Excel -> download JSON -> import to Tally*. The
generated workbook is a costing table with no GSTIN / state / country cells, so
that path dropped the buyer and seller identity entirely and the voucher went out
with ``countryofresidence``, ``placeofsupply`` and ``consigneestatename`` blank.
Tally then substitutes its own company defaults for a blank field - which is why
the bill showed **India** for a Chinese supplier and a blank place of supply.

Fixing the exporter alone cannot solve this: with no value to emit, blank is the
honest output. The workbook itself has to carry the identity, which it now does
as OOXML custom document properties - invisible in the grid, outside the
golden-tested cell layout, and preserved across Excel saves.
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from boe_converter.excel_reader import read_workbook, write_tally_names
from boe_converter.excel_writer import ExcelGenerator
from boe_converter.models import RawValue, ReviewFlagSet
from boe_converter.tally_exporter import TallyExporter

from tests.test_tally_export import make_doc, make_line


def _xlsx(doc) -> bytes:
    return ExcelGenerator(use_formulas=False).generate(doc, ReviewFlagSet([]))


def _doc():
    """A document with the identity a BOE actually carries."""
    return make_doc([make_line(1, 0.05, 100, "KGS", "9")])


def _voucher(computed, **kw) -> dict:
    return TallyExporter(**kw).build(computed, 96.05, cost_centre="CO-49 CTN-1100")[
        "tallymessage"
    ][0]


# ---------------------------------------------------------------------------
# The round-trip must not lose identity
# ---------------------------------------------------------------------------
def test_generated_workbook_carries_buyer_and_seller_identity():
    rebuilt = read_workbook(_xlsx(_doc()))
    h = rebuilt.header
    assert h.buyer_gstin.parsed == "27AAYFG7003K1ZW"
    assert h.buyer_state.parsed == "Maharashtra"
    assert h.buyer_pincode.parsed == "400080"
    assert h.seller_country.parsed == "China"
    assert h.buyer_address.parsed == "A/43, DEVIDAYAL ROAD, MULUND"
    assert h.seller_address.parsed == "ROOM 203, YIWU"


def test_bug002_003_004_correct_without_any_stored_profile():
    """The exact reported path, with no stored buyer/seller configured."""
    v = _voucher(read_workbook(_xlsx(_doc())))
    assert v["countryofresidence"] == "China"      # BUG-002
    assert v["placeofsupply"] == "Maharashtra"     # BUG-003
    assert v["consigneestatename"] == "Maharashtra"  # BUG-004
    assert v["cmpgststate"] == "Maharashtra"
    assert v["cmpgstin"] == "27AAYFG7003K1ZW"


def test_no_field_is_left_blank_for_tally_to_default():
    """A blank field is what let Tally substitute India - none may remain."""
    v = _voucher(read_workbook(_xlsx(_doc())))
    for key in (
        "countryofresidence",
        "placeofsupply",
        "consigneestatename",
        "cmpgststate",
        "cmpgstin",
    ):
        assert v[key], f"{key} is blank - Tally will substitute its own default"


# ---------------------------------------------------------------------------
# The carrier must not disturb the workbook
# ---------------------------------------------------------------------------
def test_workbook_still_has_exactly_one_sheet():
    """Req 8.1 - the metadata must not appear as an extra sheet."""
    wb = load_workbook(io.BytesIO(_xlsx(_doc())))
    assert wb.sheetnames == ["Sheet1"]


def test_identity_is_not_written_into_any_visible_cell():
    """It lives in document properties, not in the golden-tested grid."""
    ws = load_workbook(io.BytesIO(_xlsx(_doc()))).active
    seen = {
        str(c.value)
        for row in ws.iter_rows(min_row=1, max_row=12)
        for c in row
        if c.value is not None
    }
    assert "27AAYFG7003K1ZW" not in seen
    assert "China" not in seen


def test_identity_survives_the_tally_name_mapping_roundtrip():
    """Step 2 re-saves the workbook; the identity must come through intact."""
    edited = write_tally_names(_xlsx(_doc()), {1: "MAPPED NAME"})
    rebuilt = read_workbook(edited)
    assert rebuilt.header.seller_country.parsed == "China"
    assert _voucher(rebuilt)["placeofsupply"] == "Maharashtra"


# ---------------------------------------------------------------------------
# Backwards compatibility
# ---------------------------------------------------------------------------
def test_workbook_without_identity_metadata_still_reads():
    """Workbooks downloaded before this change must not break."""
    stripped = io.BytesIO()
    wb = load_workbook(io.BytesIO(_xlsx(_doc())))
    wb.custom_doc_props = type(wb.custom_doc_props)()
    wb.save(stripped)

    rebuilt = read_workbook(stripped.getvalue())
    assert rebuilt.lines  # still a usable document
    assert rebuilt.header.buyer_gstin.is_missing
    assert rebuilt.header.seller_country.is_missing


def test_absent_identity_is_still_never_invented():
    """No metadata and no stored profile -> blank, not a fabricated state."""
    doc = make_doc(
        [make_line(1, 0.05, 100, "KGS", "9")],
        buyer_gstin=RawValue.missing(),
        buyer_state=RawValue.missing(),
        seller_country=RawValue.missing(),
    )
    v = _voucher(read_workbook(_xlsx(doc)))
    assert v["placeofsupply"] == ""
    assert v["countryofresidence"] == ""
