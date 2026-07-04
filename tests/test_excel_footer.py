"""The 'customer and other expenses' footer must reflect the BOE, not the template.

The workbook is cloned from a style template that carries the sample BOE's
values in E75-E78 (B/E NO, B/E DATE, BL NO, BL DATE). These tests assert those
cells are populated from THIS document's header and cleared when the field is
missing, so the template's values never bleed through as "defaults".
"""

from __future__ import annotations

import io

from openpyxl import load_workbook

from boe_converter.excel_writer import ExcelGenerator
from boe_converter.models import (
    ComputedDocument,
    HeaderBlock,
    RawValue,
    ReviewFlagSet,
    Totals,
)


def _rv(v) -> RawValue:
    return RawValue(raw_text=str(v), parsed=v)


def _doc(**over) -> ComputedDocument:
    base = dict(
        company_name="M/S TEST CO",
        party_name=_rv("SOME SUPPLIER"),
        usd_rate=95.3,
        details=_rv("CO-99 CTN-500"),
        invoice_no=_rv("INV999"),
        invoice_date=_rv("01/01/2026"),
        be_no=_rv("9999999"),
        be_date=_rv("31/12/2025"),
        bl_no=_rv("BL-ABC-123"),
        bl_date=_rv("15/12/2025"),
        invoice_amount=_rv(100.0),
        invoice_currency=_rv("USD"),
        package_count=_rv(500),
        container_details=RawValue.missing(),
    )
    base.update(over)
    return ComputedDocument(header=HeaderBlock(**base), lines=[], totals=Totals(), flags=[])


def _ws(doc):
    xlsx = ExcelGenerator(use_formulas=False).generate(doc, ReviewFlagSet([]))
    return load_workbook(io.BytesIO(xlsx)).active


def test_footer_reflects_boe_header():
    ws = _ws(_doc())
    assert ws["E75"].value == "9999999"        # B/E NO
    assert ws["E76"].value == "31/12/2025"     # B/E DATE
    assert ws["E77"].value == "BL-ABC-123"     # BL NO
    assert ws["E78"].value == "15/12/2025"     # BL DATE


def test_footer_clears_missing_fields_no_template_bleed():
    ws = _ws(_doc(bl_no=RawValue.missing(), bl_date=RawValue.missing()))
    # Missing BL must clear the template's value rather than leave it.
    assert ws["E77"].value is None
    assert ws["E78"].value is None
    # Present fields still populate.
    assert ws["E75"].value == "9999999"
    assert ws["E76"].value == "31/12/2025"
