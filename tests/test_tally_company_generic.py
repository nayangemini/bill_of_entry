"""Tests for the per-company parts of the Tally export.

Nothing about a company's Tally can be derived from the Bill of Entry: ledger
names are hand-made inside each company's Tally, the GST registration has a name,
and the Tally company name may carry a financial-year suffix. These prove the
exporter takes all of that from stored per-company data rather than inventing it.

Real evidence driving this (CO-04 data set vs the earlier company):
  - ledger ``Igst Purchase @18.00%``  (convention would give ``IGST Purchase @ 18.00 %``)
  - party ledger ``Prayan Impex Company Limited`` vs ``DIA IMPEX COMPANY LIMITED``
  - ``basicbuyername`` ``Gemini Unicom LLP (F.Y. 2026-27)``
All data below is synthetic.
"""

from __future__ import annotations

import pytest

from boe_converter.models import RawValue
from boe_converter.tally_exporter import (
    KIND_CUSTOM_DUTY,
    KIND_IGST_PAYABLE,
    KIND_IGST_PURCHASE,
    KIND_PURCHASE,
    KIND_TAX_FREE,
    LEDGER_KINDS,
    CompanyProfile,
    LedgerBook,
    SellerProfile,
    TallyExporter,
)

from tests.test_tally_export import make_doc, make_line, rv


def _doc(*rates):
    return make_doc([make_line(i + 1, r, 100, "KGS", "9") for i, r in enumerate(rates)])


def _ledger_names(doc, **kw) -> list[str]:
    entries = TallyExporter(**kw).build(doc, 93.8)["tallymessage"][0]["allledgerentries"]
    return [e["ledgername"] for e in entries]


# ---------------------------------------------------------------------------
# LedgerBook lookup
# ---------------------------------------------------------------------------
def test_stored_name_wins_over_the_convention():
    book = LedgerBook({(KIND_IGST_PURCHASE, 1800): "Igst Purchase @18.00%"})
    assert book.resolve(KIND_IGST_PURCHASE, 0.18) == "Igst Purchase @18.00%"


def test_convention_is_the_fallback_when_nothing_is_stored():
    assert LedgerBook().resolve(KIND_IGST_PURCHASE, 0.18) == "IGST Purchase @ 18.00 %"


def test_lookup_does_not_depend_on_float_equality():
    """0.07 and 0.1+0.1-0.13 are different floats but the same 700 bp rate."""
    book = LedgerBook({(KIND_PURCHASE, 700): "Factory Purchase (Import 7%)"})
    assert book.resolve(KIND_PURCHASE, 0.1 + 0.1 - 0.13) == "Factory Purchase (Import 7%)"


def test_from_records_builds_the_book_from_store_rows():
    class Row:
        def __init__(self, kind, rate_bp, name):
            self.kind, self.rate_bp, self.name = kind, rate_bp, name

    book = LedgerBook.from_records(
        [Row(KIND_PURCHASE, 500, "Factory Purchase (Import 5%)"), Row(KIND_TAX_FREE, 0, "")]
    )
    assert book.get(KIND_PURCHASE, 0.05) == "Factory Purchase (Import 5%)"
    assert book.get(KIND_TAX_FREE, 0.0) is None  # blank names are not stored


@pytest.mark.parametrize("kind", LEDGER_KINDS)
def test_every_kind_resolves_to_something(kind):
    assert LedgerBook().resolve(kind, 0.05)


# ---------------------------------------------------------------------------
# The exporter actually uses the book
# ---------------------------------------------------------------------------
def test_voucher_uses_the_companys_real_ledger_names():
    """The exact mismatch found in the CO-04 data."""
    book = LedgerBook({(KIND_IGST_PURCHASE, 1800): "Igst Purchase @18.00%"})
    names = _ledger_names(_doc(0.18), ledgers=book)
    assert "Igst Purchase @18.00%" in names
    assert "IGST Purchase @ 18.00 %" not in names


def test_all_five_ledger_kinds_are_taken_from_the_book():
    book = LedgerBook(
        {
            (KIND_PURCHASE, 500): "P5",
            (KIND_IGST_PURCHASE, 500): "IP5",
            (KIND_IGST_PAYABLE, 500): "IY5",
            (KIND_TAX_FREE, 0): "TF",
            (KIND_CUSTOM_DUTY, 0): "CD",
        }
    )
    names = _ledger_names(_doc(0.05, 0.0), ledgers=book)
    for expected in ("P5", "IP5", "IY5", "TF", "CD"):
        assert expected in names, f"{expected} missing from {names}"


def test_required_ledger_names_reflects_the_book():
    book = LedgerBook({(KIND_CUSTOM_DUTY, 0): "Customs Duty A/c"})
    doc = _doc(0.05)
    assert "Customs Duty A/c" in TallyExporter(ledgers=book).required_ledger_names(doc)


# ---------------------------------------------------------------------------
# Derived names are reported, never silently used
# ---------------------------------------------------------------------------
def test_derived_names_are_reported_when_nothing_is_stored():
    derived = dict(
        (k, n) for k, n in TallyExporter().derived_ledger_names(_doc(0.05))
    )
    assert derived[KIND_PURCHASE] == "Factory Purchase (Import 5%)"
    assert derived[KIND_CUSTOM_DUTY] == "Custom Duty Payable"


def test_stored_names_are_not_reported_as_derived():
    book = LedgerBook(
        {
            (KIND_PURCHASE, 500): "P5",
            (KIND_IGST_PURCHASE, 500): "IP5",
            (KIND_IGST_PAYABLE, 500): "IY5",
            (KIND_CUSTOM_DUTY, 0): "CD",
        }
    )
    assert TallyExporter(ledgers=book).derived_ledger_names(_doc(0.05)) == []


def test_partially_stored_book_reports_only_the_gap():
    book = LedgerBook(
        {
            (KIND_PURCHASE, 500): "P5",
            (KIND_IGST_PURCHASE, 500): "IP5",
            (KIND_IGST_PAYABLE, 500): "IY5",
        }
    )
    kinds = [k for k, _ in TallyExporter(ledgers=book).derived_ledger_names(_doc(0.05))]
    assert kinds == [KIND_CUSTOM_DUTY]


# ---------------------------------------------------------------------------
# Party ledger: the stored name is the ledger name, verbatim
# ---------------------------------------------------------------------------
def test_stored_seller_name_is_used_verbatim_as_the_party_ledger():
    v = TallyExporter(seller=SellerProfile(name="DIA IMPEX COMPANY LIMITED")).build(
        _doc(0.05), 93.8
    )["tallymessage"][0]
    assert v["partyledgername"] == "DIA IMPEX COMPANY LIMITED"


def test_party_ledger_is_title_cased_only_without_a_stored_seller():
    v = TallyExporter().build(_doc(0.05), 93.8)["tallymessage"][0]
    assert v["partyledgername"] == "Prayan Impex Company Limited"


# ---------------------------------------------------------------------------
# Per-company identity fields
# ---------------------------------------------------------------------------
def test_tally_company_name_drives_basicbuyername():
    v = TallyExporter(
        company=CompanyProfile(tally_name="Gemini Unicom LLP (F.Y. 2026-27)")
    ).build(_doc(0.05), 93.8)["tallymessage"][0]
    assert v["basicbuyername"] == "Gemini Unicom LLP (F.Y. 2026-27)"
    # ...while the mailing name stays plain.
    assert v["consigneemailingname"] == "GEMINI UNICOM LLP"


def test_basicbuyername_defaults_to_the_boe_buyer():
    v = TallyExporter().build(_doc(0.05), 93.8)["tallymessage"][0]
    assert v["basicbuyername"] == "GEMINI UNICOM LLP"


# ---------------------------------------------------------------------------
# Voucher date is the caller's decision
# ---------------------------------------------------------------------------
def test_voucher_date_defaults_to_the_be_date():
    v = TallyExporter().build(_doc(0.05), 93.8)["tallymessage"][0]
    assert v["date"] == "20260627"          # the BE date from the shared fixture
    assert v["effectivedate"] == "20260627"


def test_voucher_date_can_be_chosen_by_the_caller():
    v = TallyExporter().build(_doc(0.05), 93.8, voucher_date="20260420")["tallymessage"][0]
    assert v["date"] == "20260420"
    assert v["effectivedate"] == "20260420"
    # the supplier invoice keeps its own date regardless
    assert v["referencedate"] == "20260601"


# ---------------------------------------------------------------------------
# A foreign workbook must be refused, not silently misread
# ---------------------------------------------------------------------------
def test_foreign_workbook_layout_is_refused_loudly():
    """A sheet whose columns are shifted would be read into the wrong fields."""
    import io as _io

    from openpyxl import Workbook

    from boe_converter.excel_reader import ExcelReadError, read_workbook

    wb = Workbook()
    ws = wb.active
    # The client's own sheet: labels one column left of ours, header on row 15.
    ws["C1"], ws["D1"] = "Company name", "Gemini Unicom LLP"
    ws["C2"], ws["D2"] = "Party Name", "PRAYAN IMPEX COMPANY LIMITED"
    for col, label in enumerate(["Sr. no.", "PARTY NAME", "AS PER TALLY NAME"], start=1):
        ws.cell(row=15, column=col, value=label)
    ws.cell(row=16, column=1, value=1)
    buf = _io.BytesIO()
    wb.save(buf)

    with pytest.raises(ExcelReadError) as exc:
        read_workbook(buf.getvalue())
    assert "D1" in str(exc.value)          # names the offending cell
    assert "Company name" in str(exc.value)  # ...and what it should have held


def test_our_own_generated_workbook_still_reads():
    from boe_converter.excel_writer import ExcelGenerator
    from boe_converter.excel_reader import read_workbook
    from boe_converter.models import ReviewFlagSet

    doc = _doc(0.05)
    rebuilt = read_workbook(ExcelGenerator(use_formulas=False).generate(doc, ReviewFlagSet([])))
    assert len(rebuilt.lines) == len(doc.lines)


# ---------------------------------------------------------------------------
# Stored per-company identity reaches the profile
# ---------------------------------------------------------------------------
def test_company_profile_carries_tax_unit_and_tally_name():
    from boe_converter.tally_exporter import company_profile_for

    class Rec:
        gstin = state = pincode = ""
        address_lines: list[str] = []
        tax_unit = "Maharashtra Registration"
        tally_name = "Gemini Unicom LLP (F.Y. 2026-27)"

    doc = make_doc(
        [make_line(1, 0.05, 100, "KGS", "9")],
        buyer_gstin=RawValue.missing(),
        buyer_state=rv("Maharashtra"),
    )
    prof = company_profile_for(doc.header, Rec())
    assert prof.tax_unit == "Maharashtra Registration"
    assert prof.tally_name == "Gemini Unicom LLP (F.Y. 2026-27)"

    v = TallyExporter(company=prof).build(doc, 93.8)["tallymessage"][0]
    assert v["vchstatustaxunit"] == "Maharashtra Registration"
    assert v["basicbuyername"] == "Gemini Unicom LLP (F.Y. 2026-27)"


# ---------------------------------------------------------------------------
# The voucher must balance EXACTLY - Tally's own vouchers sum to 0.00
# ---------------------------------------------------------------------------
from decimal import Decimal  # noqa: E402


def _exact_total(entries) -> Decimal:
    """Sum the emitted amount strings in decimal, as Tally reads them."""
    return sum(Decimal(e["amount"]) for e in entries)


@pytest.mark.parametrize(
    "rates",
    [(0.05,), (0.18,), (0.0,), (0.05, 0.18), (0.0, 0.05, 0.18), (0.025, 0.12, 0.28)],
)
def test_voucher_balances_to_exactly_zero(rates):
    entries = TallyExporter().build(_doc(*rates), 93.8)["tallymessage"][0][
        "allledgerentries"
    ]
    assert _exact_total(entries) == Decimal("0.00")


def test_purchase_ledgers_equal_the_sum_of_their_stock_items():
    """The residue must never be pushed into a ledger pinned to stock values."""
    entries = TallyExporter().build(_doc(0.05, 0.18), 93.8)["tallymessage"][0][
        "allledgerentries"
    ]
    checked = 0
    for e in entries:
        inv = e.get("inventoryallocations")
        if not inv:
            continue
        assert Decimal(e["amount"]) == sum(Decimal(i["amount"]) for i in inv)
        checked += 1
    assert checked >= 2


def test_party_bill_allocation_matches_the_party_amount():
    entries = TallyExporter().build(_doc(0.05, 0.18), 93.8)["tallymessage"][0][
        "allledgerentries"
    ]
    party = next(e for e in entries if e.get("ispartyledger"))
    assert Decimal(party["billallocations"][0]["amount"]) == Decimal(party["amount"])
